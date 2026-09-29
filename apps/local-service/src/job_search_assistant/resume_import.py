import hashlib
import re
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED_RESUME_SUFFIXES = {".pdf", ".docx", ".txt", ".md", ".markdown"}
MAX_RESUME_BYTES = 10 * 1024 * 1024


@dataclass
class ParsedResume:
    profile: dict[str, str]
    projects: list[dict[str, object]]
    resume: dict[str, object]


def parse_resume(filename: str, content: bytes) -> ParsedResume:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_RESUME_SUFFIXES:
        raise ValueError("仅支持 PDF、DOCX、TXT 和 Markdown 简历")
    if not content:
        raise ValueError("简历文件为空")
    if len(content) > MAX_RESUME_BYTES:
        raise ValueError("简历文件不能超过 10 MB")

    text = _normalize_text(_extract_text(suffix, content))
    if not text:
        raise ValueError("没有从简历中提取到可用文字")

    source_hash = hashlib.sha256(content).hexdigest()
    profile = _parse_profile(text)
    project_section = _section(
        text, ("项目经历", "项目经验", "项目实践"), ("工作经历", "教育经历")
    )
    projects = extract_projects_from_text(text, source_hash, filename)
    if project_section and not projects:
        projects.append(
            {
                "name": f"{Path(filename).stem} · 项目经历",
                "data": {
                    "summary": project_section,
                    "tags": _extract_tags(project_section),
                    "source": "resume-import",
                    "sourceFile": filename,
                    "sourceHash": source_hash,
                },
            }
        )

    profile["workExperience"] = _section(
        text, ("工作经历", "工作经验"), ("项目经历", "教育经历")
    )
    profile["education"] = _section(text, ("教育经历", "教育背景"), ())
    tech_stack = _section(
        text,
        ("技术栈", "专业技能", "技能清单", "技能"),
        ("个人优势", "工作经历", "项目经历", "教育经历"),
    )
    profile["techStack"] = tech_stack or _extract_tags(text)
    profile = {key: value for key, value in profile.items() if value}

    return ParsedResume(
        profile=profile,
        projects=projects,
        resume={
            "format": suffix.removeprefix(".").upper(),
            "notes": "通过简历文件自动导入",
            "fileName": filename,
            "rawText": text,
            "sourceHash": source_hash,
            "importedAt": datetime.now(UTC).isoformat(),
        },
    )


def _extract_text(suffix: str, content: bytes) -> str:
    if suffix == ".pdf":
        try:
            pages = PdfReader(BytesIO(content)).pages
            return "\n".join(page.extract_text() or "" for page in pages)
        except Exception as error:
            raise ValueError("PDF 解析失败，请确认文件没有加密或损坏") from error
    if suffix == ".docx":
        try:
            document = Document(BytesIO(content))
        except Exception as error:
            raise ValueError("DOCX 解析失败，请确认文件没有损坏") from error
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_cells = [
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        ]
        return "\n".join([*paragraphs, *table_cells])
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("文本文件编码无法识别")


def _normalize_text(text: str) -> str:
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.replace("\r", "").split("\n")]
    return "\n".join(line for line in lines if line)


def _parse_profile(text: str) -> dict[str, str]:
    lines = text.splitlines()
    profile: dict[str, str] = {}
    labelled_name = _match(text, r"(?:姓名|名字)[：:]\s*([^\n]{2,20})")
    if labelled_name:
        profile["displayName"] = labelled_name
    elif lines and len(lines[0]) <= 20 and not re.search(r"\d|简历|求职", lines[0]):
        profile["displayName"] = lines[0]

    email = _match(text, r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}")
    phone = _match(text, r"(?<!\d)(?:\+?86[- ]?)?1[3-9]\d{9}(?!\d)")
    target = _match(text, r"(?:目标岗位|求职意向|期望职位)[：:]\s*([^\n]+)")
    years = _match(text, r"(\d{1,2})\s*年(?:工作)?经验")
    cities = _match(text, r"(?:期望城市|工作地点|期望地点)[：:]\s*([^\n]+)")
    summary = _section(
        text,
        ("个人简介", "个人总结", "自我评价", "个人优势"),
        ("技能", "工作经历", "项目经历"),
    )
    if email:
        profile["email"] = email
    if phone:
        profile["phone"] = phone
    if target:
        profile["targetRoles"] = target
    if years:
        profile["yearsExperience"] = years
    if cities:
        profile["cities"] = cities
    if summary:
        profile["summary"] = summary
    return profile


def _section(text: str, starts: tuple[str, ...], ends: tuple[str, ...]) -> str:
    start_pattern = "|".join(re.escape(value) for value in starts)
    start = re.search(rf"(?m)^\s*(?:{start_pattern})\s*[：:]?\s*$", text)
    if not start:
        return ""
    remainder = text[start.end() :].lstrip()
    if ends:
        end_pattern = "|".join(re.escape(value) for value in ends)
        end = re.search(rf"(?m)^\s*(?:{end_pattern})\s*[：:]?\s*$", remainder)
        if end:
            remainder = remainder[: end.start()]
    return remainder.strip()


def _match(text: str, pattern: str) -> str:
    match = re.search(pattern, text, flags=re.IGNORECASE)
    return match.group(1 if match.lastindex else 0).strip() if match else ""


def extract_projects_from_text(
    text: str, source_hash: str, filename: str
) -> list[dict[str, object]]:
    """Best-effort fallback when a cloud model is unavailable or returns no projects.

    Resume project headings usually contain a project name, a role and a date range. We use
    those anchors instead of splitting by character count so every fallback record remains a
    real project entity that can receive its own database ID and vector.
    """
    heading = re.compile(
        r"(?m)^(?P<name>[^\n]{2,80}?)\s+"
        r"(?P<role>核心开发|项目负责人|负责人|核心成员|主要开发|独立开发|技术负责人)\s+"
        r"(?P<start>(?:19|20)\d{2}[./-]\d{1,2})\s*[-–—至~]+\s*"
        r"(?P<end>(?:(?:19|20)\d{2}[./-]\d{1,2}|至今|现在))\s*$"
    )
    matches = list(heading.finditer(text))
    projects: list[dict[str, object]] = []
    stop_heading = re.compile(r"(?m)^\s*(?:工作经历|教育经历|教育背景)\s*$")
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        stop = stop_heading.search(text, match.end(), end)
        if stop:
            end = stop.start()
        body = text[match.end() : end].strip()
        if not body:
            continue
        technologies = _extract_tags(body)
        projects.append(
            {
                "name": match.group("name").strip(),
                "data": {
                    "summary": body.splitlines()[0].strip(),
                    "role": match.group("role").strip(),
                    "technologies": [item for item in technologies.split(",") if item],
                    "achievements": [
                        line.strip() for line in body.splitlines()[1:] if line.strip()
                    ],
                    "tags": technologies,
                    "startDate": match.group("start").strip(),
                    "endDate": match.group("end").strip(),
                    "evidence": f"{match.group(0).strip()}\n{body}",
                    "extractionMethod": "local-heading-fallback",
                    "source": "resume-import",
                    "sourceFile": filename,
                    "sourceHash": source_hash,
                },
            }
        )
    return projects


def _extract_tags(text: str) -> str:
    known = (
        "Python",
        "FastAPI",
        "Java",
        "Go",
        "Vue",
        "React",
        "TypeScript",
        "RAG",
        "LangChain",
        "LangGraph",
        "PostgreSQL",
        "MySQL",
        "Redis",
        "Docker",
        "LLM",
        "Agent",
    )
    return ",".join(tag for tag in known if re.search(re.escape(tag), text, re.IGNORECASE))
